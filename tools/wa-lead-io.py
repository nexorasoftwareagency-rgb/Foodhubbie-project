#!/usr/bin/env python3
# read/update bridge for tools/wa-sales-agent.cjs  (xlsx <-> JSON)
#   read   -> prints JSON rows of Leads sheet (with CRM IDs, auto-assigned and persisted)
#   update -> argv[3] JSON: {"crm":"LD003","patch":{"status":"Replied","attempts":2,...}}
# CRM ID lives in a dedicated "CRM ID" column; assigns LD### to new rows, keeps existing.
import json, os, sys, re
from datetime import date, datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

XLSX = sys.argv[2]
CRM_COL_HDR = "CRM ID"

def headers(ws):
    hdrs = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=2, column=c).value
        if v: hdrs[v.strip().lower()] = c
    return hdrs

def main():
    mode = sys.argv[1]
    wb = load_workbook(XLSX)
    ws = wb["Leads"]
    hdrs = headers(ws)

    # --- ensure CRM column exists ---
    if CRM_COL_HDR.lower() not in hdrs:
        col = max(hdrs.values()) + 1
        ws.cell(row=2, column=col, value=CRM_COL_HDR)
        ws.cell(row=2, column=col).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=2, column=col).fill = PatternFill("solid", fgColor="1F2937")
        ws.column_dimensions[ws.cell(row=2, column=col).column_letter].width = 12
        hdrs[CRM_COL_HDR.lower()] = col
    crm_col = hdrs["crm id"]

    # --- assign CRM ids ---
    used = set()
    for r in range(3, ws.max_row + 1):
        v = ws.cell(row=r, column=crm_col).value
        if v:
            used.add(str(v))
    nxt = max([int(re.match(r"LD(\d+)", u).group(1)) for u in used if re.match(r"LD(\d+)", u)], default=0)
    rows = []
    for r in range(3, ws.max_row + 1):
        name = ws.cell(row=r, column=2).value
        if not name: continue
        if not ws.cell(row=r, column=crm_col).value:
            nxt += 1
            ws.cell(row=r, column=crm_col, value=f"LD{nxt:03d}")
        rec = {}
        for k in ("crm id", "restaurant name", "city / area", "rating (maps)", "phone", "whatsapp?",
                  "website", "owner name", "outlet type", "dine-in fit (tables)", "source (query)",
                  "status", "channel used", "attempts", "first contact", "last contact",
                  "reply summary", "est. ₹/month (own math)", "objections", "notes", "msg lang"):
            if k in hdrs:
                rec[k] = ws.cell(row=r, column=hdrs[k]).value
        rec["_col"] = {k: hdrs[k] for k in hdrs}
        rec["_row"] = r
        rows.append(rec)
    if mode == "read":
        wb.save(XLSX)
        out = []
        for rec in rows:
            clean = lambda v, k="": ("" if v is None else str(v).replace("\ufffd", ""))
            dmy = lambda v: ("" if v is None
                             else (v.strftime("%d/%m/%Y") if hasattr(v, "strftime") else str(v).replace("\ufffd", "")))
            out.append({
                "crm": clean(rec["crm id"]), "name": clean(rec["restaurant name"]),
                "phone": clean(rec["phone"]).replace(" ", ""), "addr": clean(rec["city / area"]),
                "rating": clean(rec["rating (maps)"]), "outletType": clean(rec["outlet type"]),
                "fit": clean(rec["dine-in fit (tables)"]), "website": clean(rec["website"]),
                "ownerName": clean(rec["owner name"]), "waKnown": clean(rec["whatsapp?"]),
                "status": clean(rec["status"]), "attempts": rec["attempts"] or 0,
                "source": clean(rec["source (query)"]),
                "lastContact": dmy(rec["last contact"]), "notes": clean(rec["notes"]),
                "lang": clean(rec["msg lang"]) if "msg lang" in rec else "",
                "instagram": clean(rec.get("instagram")),
                "replySummary": clean(rec["replySummary"]) if "replySummary" in rec else clean(rec["reply summary"]),
                "row": rec["_row"], "cols": rec["_col"],
            })
        print(json.dumps(out, ensure_ascii=True))
    elif mode == "update":
        raw = sys.argv[3]
        job = json.loads(open(raw, encoding="utf-8").read() if os.path.isfile(raw) else raw)
        hdrs2 = headers(ws)
        row = next(r["_row"] for r in rows if r.get("crm id") == job["crm"])
        for k, v in (job.get("patch") or {}).items():
            key = k.strip().lower()
            # dd/mm/yyyy strings for contact columns become real Excel dates (formulas + date math)
            if key in ("first contact", "last contact") and isinstance(v, str) and re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}", v):
                v = datetime.strptime(v, "%d/%m/%Y").date()
            if key in hdrs2:
                ws.cell(row=row, column=hdrs2[key], value=v)
            else:
                print(f"WARN no col: {k}", file=sys.stderr)
        wb.save(XLSX)
        print(json.dumps({"saved": job["crm"], "row": row}))
    else:
        print(json.dumps({"error": "mode?" }))

if __name__ == "__main__":
    main()
